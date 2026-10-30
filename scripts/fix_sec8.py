"""重写第八节：按职能分工 + 场景认领制（不写死场景）。"""
from __future__ import annotations

from pathlib import Path

from docx import Document

SRC = Path(r"C:/Users/蔡朝彬/Desktop/硅基文明_方案_v2.docx")


def set_text(par, text):
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    par.add_run(text)


def find_par(doc, prefix):
    for p in doc.paragraphs:
        if (p.text or "").strip().startswith(prefix):
            return p
    return None


def make_table(doc, rows, anchor_el):
    width = max(len(r) for r in rows)
    tbl = doc.add_table(rows=0, cols=width)
    tbl.style = "Table Grid"
    for r in rows:
        cells = tbl.add_row().cells
        for j in range(width):
            cells[j].text = r[j] if j < len(r) else ""
    anchor_el.addnext(tbl._tbl)
    return tbl


def main():
    doc = Document(str(SRC))

    p_intro = find_par(doc, "分工原则：")
    p_note = find_par(doc, "底座对队员暴露的接口")

    if p_intro is not None:
        set_text(p_intro, "分工按**职能**划分，不按场景划分：既要保证当前测试场景有人负责，"
                          "也要保证**随时能追加新场景、随时能换人**。"
                          "底座提供「插件协议 + 实验装置」，场景 owner 只写配置、跑批次、写报告，不改基座代码。")

    old = None
    for tb in doc.tables:
        if (tb.rows[0].cells[0].text or "").strip() == "成员":
            old = tb
            break
    if old is not None:
        old._tbl.getparent().remove(old._tbl)

    if p_note is not None:
        set_text(p_note, "底座对队员暴露的接口：scenarios/*.yaml（场景参数）、plugin.yaml（插件元数据）、"
                          "run_experiment.py 的 --plan 预设、fill_experiment.py 的报告模板；"
                          "新增场景不需要改基座，也不需要改本表。")
        anchor = p_note._p
        team = make_table(doc, [
            ["成员", "职能", "具体职责", "交付物"],
            ["蔡朝彬", "基座与实验装置（owner）",
             "会话/频道/回合/裁判、插件协议、世界与指标、批次与清单、报告与图表、回放台、角色卡、权力分析、前端；"
             "定义「新场景接入」规范与验收门槛",
             "基座代码与文档、插件协议、装置脚本、框架稳定性测试报告"],
            ["从文榜", "研究设计与结论", "研究问题与假设、条件矩阵设计原则、结论边界、对外汇报与答辩",
             "研究方案、最终报告"],
            ["梁俊灏", "场景接入（插件开发）",
             "按插件协议接入**自选场景**（不限现有三个）；保证其可运行并产出场景报告", "新场景插件 + 场景报告"],
            ["唐添烨", "数据与指标", "指标口径、统计方法（CI/显著性）、LLM judge 抽样、图表",
             "数据集、指标定义、图表"],
            ["王子钊", "交互与展示", "会话台/回放台/批次页的需求与验收、演示与视频、答辩材料", "可视化、演示材料"],
        ], anchor)

        p_scene = doc.add_paragraph()
        p_scene.add_run("场景认领（可空缺、可更换；新增场景不占名额）：").bold = True
        team._tbl.addnext(p_scene._p)
        make_table(doc, [
            ["场景", "定位", "当前状态", "认领人"],
            ["commons 公地困境", "测试场景（底座验证用）", "框架稳定性测试：120 局", "待认领"],
            ["rules_horror 规则怪谈", "测试场景（底座验证用）", "框架稳定性测试：6 局", "待认领"],
            ["werewolf 狼人杀", "测试场景（底座验证用）", "插件完整，真模型批次未跑", "待认领"],
            ["discussion 圆桌讨论", "基座自检", "已完成", "无需认领"],
            ["（任意新增场景）", "由成员自选", "按插件协议接入即可", "任何人可认领"],
        ], p_scene._p)

        p_flow = doc.add_paragraph()
        p_flow.add_run("新增场景的四步（不改基座、不改分工表）：").bold = True
        tbl_anchor = doc.tables[-1]._tbl
        tbl_anchor.addnext(p_flow._p)
        anchor = doc.tables[-1]._tbl
        for text in [
            "1) 新建 plugins/<id>/plugin.py 与 plugin.yaml，实现 system_text / facts / validate / setup / run / metrics / outcome；",
            "2) 在 scenarios/ 下写该场景的参数文件（字段随场景自定）；",
            "3) 在 run_experiment.py 的 PLANS 里加一个预设（场景文件 + 条件 sweep）；",
            "4) 跑 1–2 局冒烟 → 交底座验收（能跑、指标可算、日志可回放）→ 认领人产出场景报告。",
        ]:
            p = doc.add_paragraph()
            p.add_run(text)
            anchor.addnext(p._p)
            anchor = p._p
    doc.save(str(SRC))
    print("saved; tables:", len(doc.tables))


if __name__ == "__main__":
    main()