"""把 PLAN_V2.md 转成 docx（只支持本项目用到的 Markdown 子集）。"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

FENCE = chr(96) * 3
BT = chr(96)
SENT = "\u0001"


def set_cjk_font(doc, name="微软雅黑"):
    style = doc.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)
    for i, size in ((1, 18), (2, 15), (3, 12.5)):
        h = doc.styles["Heading {}".format(i)]
        h.font.name = name
        h.font.size = Pt(size)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_inline(par, text):
    text = (text or "").replace(BT, SENT)
    tokens = re.split(r"(\*\*[^*]+\*\*|" + SENT + r"[^" + SENT + r"]+" + SENT + r")", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith(SENT) and tok.endswith(SENT) and len(tok) > 2:
            run = par.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            par.add_run(tok.replace(SENT, BT))


def main():
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    lines = src.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_cjk_font(doc)
    i = 0
    n = len(lines)
    while i < n:
        s = lines[i].strip()
        if s.startswith(FENCE):
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith(FENCE):
                block.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Pt(12)
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                for r in rows:
                    cells = table.add_row().cells
                    for j in range(width):
                        txt = r[j] if j < len(r) else ""
                        cells[j].text = ""
                        add_inline(cells[j].paragraphs[0], txt)
                doc.add_paragraph()
            continue
        if s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            p = doc.add_heading(s[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif s.startswith("> "):
            p = doc.add_paragraph()
            add_inline(p, s[2:])
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, s[2:])
        elif re.match(r"^\d+\. ", s):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", s))
        elif s:
            p = doc.add_paragraph()
            add_inline(p, s)
        i += 1
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    print("saved:", dst, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables))


if __name__ == "__main__":
    main()