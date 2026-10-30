"""修正第八节：每名队员 = 长期职能 + 当前场景责任（不再全员待认领）。"""
from __future__ import annotations

from pathlib import Path

from docx import Document

SRC = Path(r"C:/Users/蔡朝彬/Desktop/硅基文明_方案_v2.docx")


def set_text(par, text):
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    par.add_run(text)


def find_par(doc, prefix):
    for p in doc.paragraphs:
        if (p.text or "").strip().startswith(prefix):
            return p
    return None


def main():
    doc = Document(str(SRC))

    p_intro = find_par(doc, "分工按职能划分")
    if p_intro is not None:
        set_text(p_intro, "分工用两条轴：**长期职能**（谁擅长什么，稳定）+ **当前场景责任**（谁现在对哪个场景负责，可轮换）。"
                          "每名队员都有一项职能，并至少负责一个当前场景；底座只提供协议与装置，不替场景做结论。")

    p_head = find_par(doc, "场景认领（可空缺")
    if p_head is not None:
        set_text(p_head, "场景责任（当前试点，可轮换、可追加；每个场景至少 1 名 owner + 1 名职能支持）：")

    target = None
    for tb in doc.tables:
        header = [(c.text or "").strip() for c in tb.rows[0].cells]
        if "认领人" in header:
            target = tb
            break

    if target is not None:
        anchor = target._tbl.getprevious()
        target._tbl.getparent().remove(target._tbl)
        rows = [
            ["场景", "定位", "场景 owner", "职能支持", "当前状态"],
            ["commons 公地困境", "测试场景（验证底座）", "梁俊灏", "唐添烨（指标与统计）",
             "框架稳定性测试已完成 120 局；owner 接手后补 C0–C3 条件"],
            ["rules_horror 规则怪谈", "测试场景（验证底座）", "唐添烨", "王子钊（展示与回放）",
             "框架稳定性测试已完成 6 局；owner 接手后扩样并设计更难的条文"],
            ["werewolf 狼人杀", "测试场景（验证底座）", "王子钊", "梁俊灏（规则与插件）",
             "插件完整、未跑真模型；owner 接手后跑首批并出场景报告"],
            ["discussion 圆桌讨论", "基座自检", "蔡朝彬（兼）", "—", "已完成，不需要单独报告"],
            ["（后续新增场景）", "由成员自选", "任何人认领（可空缺）", "按需指定",
             "照四步接入即可；无 owner 时底座只维持可运行，不产出场景报告"],
        ]
        width = 5
        tbl = doc.add_table(rows=0, cols=width)
        tbl.style = "Table Grid"
        for r in rows:
            cells = tbl.add_row().cells
            for j in range(width):
                cells[j].text = r[j]
        if anchor is not None:
            anchor.addnext(tbl._tbl)
        else:
            p_head._p.addnext(tbl._tbl)

        # 在场景表后补三条规则
        rules = [
            "规则一：每个场景至少有一名 owner，owner 对该场景的配置、跑批与报告负责；同时至少配一名职能支持，避免单点。",
            "规则二：owner 可更换、可新增场景、也可以两人共同认领；换人只需改本表，不影响代码与其他人的进度。",
            "规则三：一个场景若长期没有 owner，底座只保证它能跑（冒烟级），不产出场景结论——这是有意的取舍，避免无人负责却挂着结论。",
        ]
        cursor = tbl._tbl
        for text in rules:
            p = doc.add_paragraph()
            p.add_run(text)
            cursor.addnext(p._p)
            cursor = p._p

    doc.save(str(SRC))
    print("saved; tables:", len(doc.tables))


if __name__ == "__main__":
    main()