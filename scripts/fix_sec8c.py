"""调整第八节：梁俊灏改为插件总舵；场景按人摊开。"""
from __future__ import annotations

from pathlib import Path

from docx import Document

SRC = Path(r"C:/Users/蔡朝彬/Desktop/硅基文明_方案_v2.docx")


def set_cell(cell, text):
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    par = cell.paragraphs[0]
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    par.add_run(text)


def main():
    doc = Document(str(SRC))
    n = 0

    # 1) 职能表：梁俊灏 -> 插件总舵
    for tb in doc.tables:
        header = [(c.text or "").strip() for c in tb.rows[0].cells]
        if "职能" not in header:
            continue
        for row in tb.rows[1:]:
            if (row.cells[0].text or "").strip() == "梁俊灏":
                set_cell(row.cells[1], "插件总舵（插件协议与评审）")
                set_cell(row.cells[2], "定义并维护插件协议与验收标准；评审所有场景插件；"
                                      "帮队员把自己想做的场景接进来；不单独背某个场景的报告")
                set_cell(row.cells[3], "插件协议、插件评审记录、接入支持")
                n += 1
        break

    # 2) 场景责任表：摊到三个人，各配支持
    for tb in doc.tables:
        header = [(c.text or "").strip() for c in tb.rows[0].cells]
        if "场景 owner" not in header:
            continue
        plan = {
            "commons": ("唐添烨", "蔡朝彬（快速冒烟）"),
            "rules_horror": ("王子钊", "唐添烨（指标口径）"),
            "werewolf": ("从文榜", "唐添烨（数据）"),
        }
        for row in tb.rows[1:]:
            key = (row.cells[0].text or "").strip().split(" ")[0]
            if key in plan:
                owner, support = plan[key]
                set_cell(row.cells[2], owner)
                set_cell(row.cells[3], support)
                n += 1
            elif key.startswith("discussion"):
                set_cell(row.cells[2], "蔡朝彬（兼）")
                n += 1
        anchor = tb._tbl
        break
    else:
        anchor = None

    if anchor is not None:
        note = doc.add_paragraph()
        note.add_run("所有场景的插件实现由梁俊灏（插件总舵）统一评审后合入；"
                     "他只是评审与规范，不承担某个场景的跑批与报告。").italic = True
        anchor.addnext(note._p)

    doc.save(str(SRC))
    print("patches:", n)


if __name__ == "__main__":
    main()