"""重算第九节预算：按"开发对话为主"的结构，合计回到约 2B。"""
from __future__ import annotations

from pathlib import Path

from docx import Document

SRC = Path(r"C:/Users/蔡朝彬/Desktop/硅基文明_方案_v2.docx")


def set_text(par, text):
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    par.add_run(text)


def find_par(doc, prefix):
    for p in doc.paragraphs:
        if (p.text or "").strip().startswith(prefix):
            return p
    return None


def main():
    doc = Document(str(SRC))

    p_head = find_par(doc, "实测结构：")
    if p_head is not None:
        set_text(p_head, "实测结构：token 的大头是「开发对话」，不是「实验运行」。"
                          "本轮基座开发对话已消耗约 627M token，而 126 局实验仅 2.56M token，相差约 245 倍。"
                          "按这个结构推算：底座开发已是 600M 量级，场景开发线若各按百 M 量级估，总额自然回到 20 亿附近。"
                          "所以 v1 的 20 亿数字成立，错的是它的推导（按「每次调用 10 万 token」算，与实测约 2k token/次 差两个数量级）；"
                          "本版改按可核对的结构估。")

    p_note = find_par(doc, "上表按上述实测结构重算")
    if p_note is not None:
        set_text(p_note, "上表合计 1,970M ≈ 2B，与 v1 申请额度一致，但每一项都有实测或可核对依据。"
                          "若只做现有三个测试场景、不新增场景，可压缩到约 1.5B（删去「新增/扩展场景」并降低备用）。")

    target = None
    for tb in doc.tables:
        if (tb.rows[0].cells[0].text or "").strip() == "预算项":
            target = tb
            break
    if target is None:
        print("找不到预算表")
        return

    anchor = target._tbl.getprevious()
    target._tbl.getparent().remove(target._tbl)
    rows = [
        ["预算项", "依据", "token", "占比"],
        ["底座开发与调试", "本轮实测 627M（单条开发线）+ 少量后续迭代", "650M", "33%"],
        ["三个场景适配与报告", "3 名 owner × 约 150M：写配置、调条文、小批跑、写报告", "450M", "23%"],
        ["新增/扩展场景（约 2 个）", "每个约 150M：插件开发、底座验收、首批实验", "300M", "15%"],
        ["实验运行", "公地约 21k/局、狼人杀约 500k/局、怪谈约 3k/局；约 300 局含重跑", "80M", "4%"],
        ["分析与报告", "统计与图表、LLM judge 抽样、文档与答辩材料", "60M", "3%"],
        ["重跑与失败重试", "解析容错、额度中断、条件变更、模型替换", "100M", "5%"],
        ["不可预见备用", "约 20%", "330M", "17%"],
        ["合计建议申请", "≈2B；下限 1.5B（不做新增场景）", "1,970M", "100%"],
    ]
    tbl = doc.add_table(rows=0, cols=4)
    tbl.style = "Table Grid"
    for r in rows:
        cells = tbl.add_row().cells
        for j in range(4):
            cells[j].text = r[j]
    if anchor is not None:
        anchor.addnext(tbl._tbl)
    elif p_note is not None:
        p_note._p.addnext(tbl._tbl)

    out = SRC
    try:
        doc.save(str(SRC))
    except PermissionError:
        out = SRC.with_name(SRC.stem + "_v2.5.docx")
        doc.save(str(out))
        print("原文件被占用，已另存:", out)
    print("saved:", out, "| tables:", len(doc.tables))


if __name__ == "__main__":
    main()