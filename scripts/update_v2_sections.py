"""在用户手改版 docx 上做定点修改：第五节边界说明、第八节团队分工、第九节 token 预算。"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document

SRC = Path(r"C:/Users/蔡朝彬/Desktop/硅基文明_方案_v2.docx")
BAK = Path(r"C:/Users/蔡朝彬/Desktop/硅基文明_方案_v2_用户修改版备份.docx")


def set_text(par, text):
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    par.add_run(text)


def find_par(doc, prefix):
    for p in doc.paragraphs:
        if (p.text or "").strip().startswith(prefix):
            return p
    return None


def add_table_after(doc, anchor_par, rows):
    width = max(len(r) for r in rows)
    tbl = doc.add_table(rows=0, cols=width)
    tbl.style = "Table Grid"
    for r in rows:
        cells = tbl.add_row().cells
        for j in range(width):
            cells[j].text = r[j] if j < len(r) else ""
    anchor_par._p.addnext(tbl._tbl)
    return tbl


def main():
    if not BAK.exists():
        shutil.copy(SRC, BAK)
        print("备份:", BAK.name)

    doc = Document(str(SRC))

    p5 = find_par(doc, "说明：本章是")
    if p5 is not None:
        set_text(p5, "说明：本章是阶段性结果，用来证明装置可用、指标可算、成本可控；方案本身不依赖这些数字成立。"
                     "这些对局由基座开发者执行，属于**框架稳定性测试**，不作为场景结论；场景结论由各场景 owner 产出。")

    p8_intro = find_par(doc, "保留 v1 的五人分工")
    p8_note = find_par(doc, "注：技术侧实际发生的工作量")
    if p8_intro is not None:
        set_text(p8_intro, "分工原则：**底座由一人负责，场景交给队员各自认领并产出报告**。"
                           "底座提供「插件协议 + 实验装置」，场景 owner 只写配置、跑批次、写报告，不改基座代码。")
    if p8_note is not None:
        set_text(p8_note, "底座对队员暴露的接口：scenarios/*.yaml（场景参数）、plugin.yaml（插件元数据）、"
                          "run_experiment.py 的 --plan 预设、fill_experiment.py 的报告模板；需要新动作时走插件的 validate/metrics 扩展点。")
        add_table_after(doc, p8_note, [
            ["成员", "角色", "交付物", "状态"],
            ["蔡朝彬", "基座与实验装置（唯一 owner）",
             "会话/频道/回合/裁判、插件协议、世界与指标、批次与清单、报告与图表、回放台、角色卡、权力分析、前端",
             "已完成；本文档 120+6 局均为其框架稳定性测试"],
            ["从文榜", "总负责人 / 研究设计", "研究问题与假设、结论边界、对外汇报、验收各场景报告", "进行中"],
            ["梁俊灏", "公地困境场景 owner", "场景配置与条件矩阵、跑批、场景报告（含失败回放）", "待启动（底座已可用）"],
            ["唐添烨", "规则怪谈场景 owner", "场景条文设计、跑批、场景报告", "待启动"],
            ["王子钊", "狼人杀场景 owner + 展示", "场景配置、跑批、场景报告 + 可视化与演示", "待启动"],
        ])

    p9_head = find_par(doc, "实测基线：120 局")
    p9_note = find_par(doc, "v1 的 20 亿 token 预算")
    if p9_head is not None:
        set_text(p9_head, "**实测结构：token 的大头是「开发对话」，不是「实验运行」。**"
                          "本轮基座开发对话已消耗约 627M token，而 126 局实验仅 2.56M token，相差约 245 倍。"
                          "因此预算应按「开发为主、实验为辅」估算，而不是按「每次调用 10 万 token」推算。")
    if p9_note is not None:
        set_text(p9_note, "下表按上述实测结构重算；v1 的 20 亿来自错误的单次调用假设，已废弃。")

    target = None
    for tb in doc.tables:
        if (tb.rows[0].cells[0].text or "").strip() == "预算项":
            target = tb
            break
    if target is not None:
        anchor = target._tbl.getprevious()
        target._tbl.getparent().remove(target._tbl)
        rows = [
            ["预算项", "依据", "token", "占比"],
            ["底座开发与调试", "本轮实测：基座 + 4 插件 + 实验装置 + 前端 + 文档", "630M", "52%"],
            ["场景适配（3 名 owner）", "每人约 60M：写配置、调条文、跑小批、写报告", "180M", "15%"],
            ["实验运行", "公地约 21k/局；狼人杀约 500k/局；合计约 240 局", "60M", "5%"],
            ["分析与报告", "统计、图表、LLM judge 抽样、文档", "40M", "3%"],
            ["重跑与失败重试", "解析容错、额度中断、条件变更", "80M", "7%"],
            ["不可预见备用", "约 20%", "210M", "18%"],
            ["合计建议申请", "含约 10% 余量", "1.2B", "100%"],
        ]
        tbl = doc.add_table(rows=0, cols=4)
        tbl.style = "Table Grid"
        for r in rows:
            cells = tbl.add_row().cells
            for j in range(4):
                cells[j].text = r[j]
        if anchor is not None:
            anchor.addnext(tbl._tbl)
        elif p9_note is not None:
            p9_note._p.addnext(tbl._tbl)

    doc.save(str(SRC))
    print("saved; tables:", len(doc.tables))


if __name__ == "__main__":
    main()